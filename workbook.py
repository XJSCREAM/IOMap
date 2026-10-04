from collections import defaultdict
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

FIELDS=[('id','记录ID / Record ID'),('io_type','IO类型 / IO Type'),('address','IO地址 / IO Address'),('module','模块/板卡 / Module/Board'),('model','模块型号 / Module Model'),('pin','Pin / 引脚'),('device','DEVICE现场设备编号 / Field Device Tag'),('original','原文描述 / Original Description'),('translated','中文描述 / Chinese Description'),('usage','使用状态 / Usage Status'),('drawing_page','图纸页号 / Drawing Page'),('pdf_page','PDF页码 / PDF Page'),('reference','电路图引用 / Circuit Reference'),('raw_address','地址提取原文 / Raw Extracted Address'),('method','识别来源 / Extraction Method'),('status','核对状态 / Review Status'),('notes','备注/待核对原因 / Notes/Review Reasons')]

def export_workbook(document):
    w=Workbook();s=w.active;s.title='IO总表';s.append([v for _,v in FIELDS])
    io_records=[r for r in document['records'] if r.get('address') or r.get('method') in ['PDF表格','PDF通道列','人工补充']]
    for r in io_records:s.append([r.get(k,'') for k,_ in FIELDS])
    dv=DataValidation(type='list',formula1='"待核对,已确认,已修正"');s.add_data_validation(dv)
    if s.max_row>1:dv.add(f'P2:P{s.max_row}')
    issues=w.create_sheet('待核对项');issues.append(['记录ID / Record ID','模块 / Module','IO地址 / IO Address','PDF页码 / PDF Page','原因 / Reason','状态 / Status'])
    for r in document['records']:
        if r['status']=='待核对':issues.append([r['id'],r['module'],r['address'],r['pdf_page'],r['notes'],'待核对'])
    for n,e in document.get('errors',{}).items():issues.append(['','','',int(n),e,'页面提取失败'])
    for n in document.get('zero_pages',[]):issues.append(['','','',n,'本页未检出候选，请对照原图检查遗漏','页面待检查'])
    groups=defaultdict(list)
    for r in io_records:groups[r['module'] or '未关联模块'].append(r)
    summary=w.create_sheet('模块汇总');summary.append(['模块 / Module','记录数（含端子） / Record Count','不同非空地址数 / Distinct Addresses','已确认使用地址数 / Confirmed Used','备用地址数 / Spare','待核对记录数 / Pending Review'])
    for module,rr in groups.items():summary.append([module,len(rr),len({r['address'] for r in rr if r['address']}),len({r['address'] for r in rr if r['address'] and r['usage']=='使用中' and r['status']!='待核对'}),len({r['address'] for r in rr if r['address'] and r['usage']=='备用'}),sum(r['status']=='待核对' for r in rr)])
    for sheet in w:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions;sheet.row_dimensions[1].height=40
        for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24556B');c.alignment=Alignment(wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                c.alignment=Alignment(wrap_text=True,vertical='top')
                if isinstance(c.value,str):c.data_type='s'
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=24
    for col in ['H','I','Q']:s.column_dimensions[col].width=50
    out=BytesIO();w.save(out);out.seek(0);return out
