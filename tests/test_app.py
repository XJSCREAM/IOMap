import io
import tempfile
import time
import unittest
import shutil
from unittest.mock import patch
from pathlib import Path
import fitz
from PIL import Image,ImageDraw,ImageFont
from openpyxl import load_workbook
import app as module
from extraction import extract_page

class AppTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.old=module.DATA;module.DATA=Path(self.tmp.name)
        self.client=module.app.test_client()
    def tearDown(self):module.DATA=self.old;self.tmp.cleanup()
    def upload(self,pdf):
        r=self.client.post('/api/upload',data={'file':(io.BytesIO(pdf),'sample.pdf')})
        self.assertEqual(r.status_code,202);key=r.json['id']
        for _ in range(100):
            d=self.client.get('/api/documents/'+key).json
            if d['state']!='processing':return key,d
            time.sleep(.05)
        self.fail('Processing did not finish')
    def test_upload_review_persistence_export(self):
        pdf=fitz.open();page=pdf.new_page();page.insert_text((80,120),'I11_00 sensor -80BQ1')
        key,d=self.upload(pdf.tobytes());self.assertEqual(d['state'],'complete')
        r=self.client.post(f'/api/documents/{key}/records',json={'pdf_page':1}).json
        changes={'address':'I11_00','device':'-80BQ1','original':'=SUM(1,2)','translated':'传感器','status':'已修正','usage':'使用中','module':'SLOT1'}
        self.assertEqual(self.client.patch(f'/api/documents/{key}/records/{r["id"]}',json=changes).status_code,200)
        saved=self.client.get(f'/api/documents/{key}').json
        self.assertEqual(saved['records'][-1]['device'],'-80BQ1')
        result=self.client.get(f'/api/documents/{key}/export');self.assertEqual(result.status_code,200)
        w=load_workbook(io.BytesIO(result.data));self.assertEqual(w.sheetnames,['IO总表','待核对项','模块汇总'])
        cell=w['IO总表'].cell(w['IO总表'].max_row,8);self.assertEqual(cell.data_type,'s')
        self.assertEqual(self.client.get(f'/api/documents/{key}/pages/1').mimetype,'image/png')
        self.assertEqual(self.client.get(f'/api/documents/{key}/pages/2').status_code,404)
        with patch.object(module,'LocalTranslator',side_effect=RuntimeError('model missing')):
            self.assertEqual(self.client.post(f'/api/documents/{key}/translate').status_code,503)
        self.assertEqual(self.client.delete(f'/api/documents/{key}/records/{r["id"]}').status_code,200)
    def test_reject_bad_file(self):
        self.assertEqual(self.client.post('/api/upload',data={'file':(io.BytesIO(b'bad'),'fake.pdf')}).status_code,400)
    @unittest.skipUnless(shutil.which('tesseract') and Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf').exists(), 'OCR integration test requires Tesseract and the Linux fixture font')
    def test_real_scanned_pdf_ocr(self):
        image=Image.new('RGB',(1400,500),'white');draw=ImageDraw.Draw(image)
        font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',48)
        draw.text((80,180),'Q11_00   conveyor safety sensor',fill='black',font=font)
        buf=io.BytesIO();image.save(buf,format='PNG')
        pdf=fitz.open();page=pdf.new_page(width=700,height=250);page.insert_image(page.rect,stream=buf.getvalue())
        rows,mode=extract_page(page,1)
        self.assertEqual(mode,'ocr');self.assertTrue(any(r['address']=='Q11_00' for r in rows),rows)

if __name__=='__main__':unittest.main()
