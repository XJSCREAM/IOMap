import unittest
import fitz
from extraction import extract_page,enrich_records
from workbook import export_workbook
from openpyxl import load_workbook

class DiagramTests(unittest.TestCase):
    def test_column_description_pin_device_and_table_evidence(self):
        pdf=fitz.open();p=pdf.new_page(width=1132,height=770)
        for x,y,t in [(54,55,'-SLOT3'),(175,440,'I15_00'),(290,440,'I15_01'),(190,425,'11'),(305,425,'21'),(145,625,'gluer ready'),(266,625,'reserve'),(560,510,'TYPE TM5SDI12D'),(1080,748,'84')]:p.insert_text((x,y),t,fontsize=7)
        p.insert_text((181,180),'/ 218B2:4 218.2172',fontsize=7,rotate=90)
        rows,mode=extract_page(p,59)
        self.assertEqual(len(rows),2)
        first=next(r for r in rows if r['address']=='I15_00')
        self.assertEqual(first['module'],'SLOT3')
        self.assertEqual(first['pin'],'11')
        self.assertNotIn('I15_01',first['original'])
        self.assertEqual(first['device'],'-218B2')
        source={'pin':'11','original':'segnale incollatore pronto\ngluer ready signal','model':'TM5SDI12D','pdf_page':168,'drawing_page':'2517','reference':'84.831'}
        enrich_records(rows,{('SLOT3','I15_00'):[source]})
        self.assertEqual(first['original'],source['original'])
        self.assertIn('PDF第168页',first['notes'])
        self.assertEqual(first['drawing_page'],'84')
        source['pin']='99';first['original']='preserve me'
        enrich_records([first],{('SLOT3','I15_00'):[source]})
        self.assertEqual(first['original'],'preserve me')
        self.assertIn('冲突',first['notes'])
    def test_device_candidate_not_exported_as_io(self):
        pdf=fitz.open();p=pdf.new_page()
        p.insert_text((100,120),'-41W1')
        rows,_=extract_page(p,1)
        self.assertEqual(len(rows),1);self.assertEqual(rows[0]['original'],'')
        rows[0]['id']=1
        w=load_workbook(export_workbook({'records':rows}))
        self.assertEqual(w['IO总表'].max_row,1)
        self.assertEqual(w['待核对项'].max_row,2)
    def test_analog_pin_is_not_plc_address(self):
        pdf=fitz.open();p=pdf.new_page(width=1132,height=770)
        for x,y,t in [(54,55,'-SLOT22'),(137,117,'AO5.11'),(171,117,'AO5.12'),(152,133,'AQ0 (I)'),(185,133,'AQ0 (U)'),(165,620,'reserve')]:p.insert_text((x,y),t,fontsize=7)
        rows,_=extract_page(p,93)
        self.assertEqual({r['pin'] for r in rows},{'AO5.11','AO5.12'})
        self.assertTrue(all(not r['address'] for r in rows))
        for i,r in enumerate(rows):r['id']=i+1
        w=load_workbook(export_workbook({'records':rows}))
        self.assertEqual(w['IO总表'].max_row,3)

if __name__=='__main__':unittest.main()
