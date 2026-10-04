import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from translate_excel import english_description, translate_workbook

class StubTranslator:
    def translate(self, text):
        if text == 'broken':
            raise RuntimeError('test failure')
        return '纸盒输送安全传感器'

class TranslationTests(unittest.TestCase):
    def test_language_boundary(self):
        self.assertEqual(english_description('scorta\nreserve'), 'reserve')
        with self.assertRaises(ValueError):
            english_description('a\nb\nc')

    def test_workbook_preservation_and_failures(self):
        with tempfile.TemporaryDirectory() as temp:
            source, output = Path(temp)/'in.xlsx', Path(temp)/'out.xlsx'
            w=Workbook(); s=w.active; s.title='IO总表'
            s.append(['记录ID', '原文描述', '中文描述', '备注/待核对原因', 'IO地址'])
            s.append([1,'italian\ncartons conveyor safety sensor',None,'中文描述待翻译核对','I11_00'])
            s.append([2,'scorta\nreserve','人工确认译文','','I11_01'])
            s.append([3,'broken',None,'','I11_02'])
            issues=w.create_sheet('待核对项')
            issues.append(['记录ID / Record ID','模块','IO地址','PDF页码','原因 / Reason','状态'])
            issues.append([1,'SLOT1','I11_00',59,'待翻译','待核对'])
            w.save(source)
            stats=translate_workbook(source,output,StubTranslator())
            self.assertEqual(stats,{'translated':1,'preserved':1,'empty':0,'failed':1})
            result=load_workbook(output)['IO总表']
            self.assertEqual(result['E2'].value,'I11_00')
            self.assertEqual(result['B2'].value,s['B2'].value)
            self.assertEqual(result['C3'].value,'人工确认译文')
            self.assertIsNone(result['C4'].value)
            self.assertIn('翻译失败',result['D4'].value)
            self.assertIn('已本地翻译',load_workbook(output)['待核对项']['E2'].value)
            with self.assertRaises(ValueError):
                translate_workbook(source,output,StubTranslator())

if __name__=='__main__':
    unittest.main()
