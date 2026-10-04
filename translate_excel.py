"""Offline translation of IOMap workbooks using an Argos/CTranslate2 model."""
import argparse
import hashlib
import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parent

class LocalTranslator:
    def __init__(self, model, cache):
        import ctranslate2
        import sentencepiece
        model = Path(model)
        if not (model / 'model/model.bin').exists():
            raise RuntimeError('本地模型未安装，请先运行 download_model.py。')
        self.engine = ctranslate2.Translator(str(model / 'model'), device='cpu', compute_type='int8', inter_threads=1, intra_threads=2)
        self.tokens = sentencepiece.SentencePieceProcessor(model_file=str(model / 'sentencepiece.model'))
        self.path = Path(cache)
        self.cache = json.loads(self.path.read_text()) if self.path.exists() else {}
        digest = hashlib.sha256((model / 'model/model.bin').read_bytes()).hexdigest()
        self.prefix = digest + ':en:zh:'
        self.glossary = json.loads((ROOT / 'glossary.en-zh.json').read_text())

    def translate(self, text):
        text = ' '.join(text.split())
        if text.lower() in self.glossary:
            return self.glossary[text.lower()]
        key = self.prefix + text
        if key not in self.cache:
            pieces = self.tokens.encode(text, out_type=str)
            result = self.engine.translate_batch([pieces], beam_size=4, max_decoding_length=256)[0]
            self.cache[key] = self.tokens.decode(result.hypotheses[0])
            self.path.parent.mkdir(parents=True, exist_ok=True)
            tmp = self.path.with_suffix('.tmp')
            tmp.write_text(json.dumps(self.cache, ensure_ascii=False, indent=2))
            tmp.replace(self.path)
        return self.cache[key]

def english_description(original):
    # This sample's bilingual cells contain Italian first, English second.
    # Explicit --source-layout en handles other workbooks; no language guessing.
    lines = [x.strip() for x in original.splitlines() if x.strip()]
    if len(lines) == 2:
        return lines[1]
    if len(lines) == 1:
        return lines[0]
    raise ValueError('多行描述的语言边界不明确，需要手动确认')

def translate_workbook(source, output, translator, layout='it-en'):
    from openpyxl import load_workbook
    source, output = Path(source), Path(output)
    if source.resolve() == output.resolve() or output.exists():
        raise ValueError('请选择新的输出路径，避免覆盖原文件或人工修改')
    workbook = load_workbook(source)
    sheet = workbook['IO总表']
    headers = [str(c.value).split(' / ')[0] for c in sheet[1]]
    required = ['原文描述', '中文描述', '备注/待核对原因']
    columns = {name: headers.index(name) + 1 for name in required}
    stats = {'translated': 0, 'preserved': 0, 'empty': 0, 'failed': 0}
    for row in range(2, sheet.max_row + 1):
        original = sheet.cell(row, columns['原文描述']).value
        cell = sheet.cell(row, columns['中文描述'])
        if cell.value:
            stats['preserved'] += 1
            continue
        if not original:
            stats['empty'] += 1
            continue
        notes = sheet.cell(row, columns['备注/待核对原因'])
        try:
            text = english_description(str(original)) if layout == 'it-en' else str(original)
            translated = translator.translate(text)
            if not translated.strip():
                raise ValueError('模型返回空译文')
            cell.value = translated
            # Treat model output as literal spreadsheet text, never as a formula.
            cell.data_type = 's'
            notes.value = str(notes.value or '').replace('中文描述待翻译核对', '中文描述已本地翻译，待人工核对')
            if '中文描述已本地翻译' not in notes.value:
                notes.value += '；中文描述已本地翻译，待人工核对'
            stats['translated'] += 1
        except Exception as exc:
            notes.value = str(notes.value or '') + '；翻译失败：' + str(exc)
            stats['failed'] += 1
    # Synchronize generated issue descriptions while retaining human status entries.
    if '待核对项' in workbook:
        notes_by_id = {sheet.cell(r, 1).value: sheet.cell(r, columns['备注/待核对原因']).value for r in range(2, sheet.max_row + 1)}
        issues = workbook['待核对项']
        reason_column = next((i for i,c in enumerate(issues[1]) if str(c.value).split(' / ')[0] in ['核对原因','原因']),None)
        if reason_column is not None:
            for row in issues.iter_rows(min_row=2):
                if row[0].value in notes_by_id:
                    row[reason_column].value = notes_by_id[row[0].value]
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return stats

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('source', type=Path)
    parser.add_argument('output', type=Path)
    parser.add_argument('--model', type=Path, default=ROOT / 'models/en_zh')
    parser.add_argument('--cache', type=Path, default=ROOT / '.cache/translations.json')
    parser.add_argument('--source-layout', choices=['it-en', 'en'], default='it-en')
    args = parser.parse_args()
    translator = LocalTranslator(args.model, args.cache)
    stats = translate_workbook(args.source, args.output, translator, args.source_layout)
    print(json.dumps(stats, ensure_ascii=False))
    if stats['failed']:
        raise SystemExit(2)

if __name__ == '__main__':
    main()
